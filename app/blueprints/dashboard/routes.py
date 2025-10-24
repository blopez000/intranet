# app/blueprints/dashboard/routes.py

from __future__ import annotations

from datetime import date as _date, timedelta as _timedelta
from datetime import date, datetime, timedelta
from calendar import monthrange
from io import BytesIO

import pandas as pd
from flask import (
    render_template, request, send_file,
    redirect, url_for, jsonify, abort, flash
)
from flask_login import login_required, current_user
from sqlalchemy import text, func, and_, or_
from sqlalchemy.sql import bindparam

from app.extensions import db
from app.blueprints.auth.routes import nivel_requerido
from app.models import Desvinculacion
from . import bp  # blueprint definido en __init__.py
from sqlalchemy import and_, or_, func


# ================= Helpers de acceso (recintos y cuentas) =================

def _user_is_super() -> bool:
    """Admin o nivel 1 (puede ver todo)."""
    try:
        return bool(current_user.is_authenticated and current_user.is_admin_or_level1())
    except Exception:
        return False


def get_allowed_recintos_obra_ids(uid: int):
    """
    OBRA_ID visibles para el usuario, derivados de sus CUENTAS asignadas.
    (user_cuentas -> recinto_cuentas -> recintos.code)
    """
    sql = text("""
        SELECT DISTINCT CAST(r.code AS UNSIGNED) AS obra_id
        FROM user_cuentas uc
        JOIN recinto_cuentas rc ON rc.cuenta_id = uc.cuenta_id AND rc.is_active = 1
        JOIN recintos r        ON r.id = rc.recinto_id
        WHERE uc.user_id = :uid AND uc.is_active = 1
    """)
    with db.engine.begin() as conn:
        return [int(row[0]) for row in conn.execute(sql, {"uid": uid}).all()]



def enforce_rid_allowed(rid: int | None):
    if _user_is_super() or not rid:
        return
    allowed = set(get_allowed_recintos_obra_ids(current_user.id))
    if rid not in allowed:
        abort(403)



def _sql_in_clause_text(col_name: str, allowed):
    """
    - allowed=None  => '' (sin filtro)
    - allowed=set() => ' AND 1=0 ' (vacío)
    - allowed={…}   => IN :rids (expanding)
    """
    if allowed is None:
        return "", {}
    if not allowed:
        return " AND 1=0 ", {}
    _ = bindparam("rids", expanding=True)
    return f" AND {col_name} IN :rids ", {"rids": list(allowed)}


def _allowed_recinto_ids():
    """
    OBRA_ID visibles para el usuario actual:
      - None  => admin/nivel1 (sin filtro)
      - set() => sin acceso
    """
    if not current_user.is_authenticated:
        return set()
    try:
        if current_user.is_admin_or_level1():
            return None
    except Exception:
        pass

    rows = db.session.execute(text("""
        SELECT DISTINCT CAST(r.code AS UNSIGNED) AS obra_id
        FROM user_cuentas uc
        JOIN recinto_cuentas rc ON rc.cuenta_id = uc.cuenta_id AND rc.is_active = 1
        JOIN recintos r        ON r.id = rc.recinto_id
        WHERE uc.user_id = :uid AND uc.is_active = 1
    """), {"uid": current_user.id}).fetchall()

    ids = {int(r[0]) for r in rows}
    return ids if ids else set()



# ===== permisos por CUENTAS / ÁREAS (derivadas de recinto_cuentas + cuentas) =====

def _allowed_cuentas(user_id: int, recinto_ids):
    """
    Devuelve:
      - None  => sin restricción por cuentas (admin/nivel 1)
      - {}    => sin acceso (no hay cuentas)
      - dict{recinto_id(obra_id): set(cuenta_code)}
    """
    if recinto_ids is None:   # admin
        return None
    if recinto_ids == set():  # sin acceso
        return {}

    rows = db.session.execute(text("""
        SELECT
          CAST(r.code AS UNSIGNED) AS recinto_id,
          c.code                    AS cuenta_code
        FROM user_cuentas uc
        JOIN cuentas c          ON c.id = uc.cuenta_id AND c.is_active = 1
        JOIN recinto_cuentas rc ON rc.cuenta_id = uc.cuenta_id AND rc.is_active = 1
        JOIN recintos r         ON r.id = rc.recinto_id
        WHERE uc.user_id = :uid AND uc.is_active = 1
    """), {"uid": user_id}).mappings().all()

    per: dict[int, set[str]] = {}
    for r in rows:
        rid = int(r["recinto_id"])
        if recinto_ids and rid not in recinto_ids:
            continue
        per.setdefault(rid, set()).add(r["cuenta_code"])

    return per if any(per.values()) else {}


def _clause_cuentas(col_recinto: str, col_cta: str, per):
    """
    Construye cláusula para pares (recinto, cuenta) compatible con MySQL + text().
    - None  -> sin filtro
    - {}    -> AND 1=0
    - dict  -> AND ( (rec=:rid_0 AND cta IN (:cta_0_0,:cta_0_1,...)) OR ... )
    """
    if per is None:
        return "", {}
    if per == {}:
        return " AND 1=0 ", {}

    parts = []
    params = {}
    idx = 0

    for rid, cuentas in per.items():
        if not cuentas:
            continue

        # Creamos un placeholder por cada cuenta para evitar "expanding" con text()
        cta_placeholders = []
        for j, cta in enumerate(sorted(cuentas)):
            pname = f"cta_{idx}_{j}"
            cta_placeholders.append(f":{pname}")
            params[pname] = cta

        pr = f"rid_{idx}"
        params[pr] = rid

        parts.append(
            f"({col_recinto} = :{pr} AND {col_cta} IN ({', '.join(cta_placeholders)}))"
        )
        idx += 1

    if not parts:
        return "", {}

    return " AND ( " + " OR ".join(parts) + " ) ", params



 # --- NUEVO: set plano de cuentas permitidas (a partir de recintos asignados) ---
def _allowed_cuentas_flat_for_current_user():
    """
    Devuelve:
      - None     -> admin / nivel 1 (sin restricción)
      - set()    -> usuario sin permisos (no verá nada)
      - set{...} -> cuentas visibles (centro_costo_area)
    """
    recinto_ids = _allowed_recinto_ids()  # None / set()
    per = _allowed_cuentas(current_user.id, recinto_ids)
    if per is None:
        return None
    if not per:
        return set()
    s = set()
    for cuentas in per.values():
        s.update(cuentas or [])
    return s




# =================== SQL base comunes ===================

SQL_INASISTENCIAS_BASE = """
SELECT DISTINCT
  CONCAT(LEFT(i.DNI, LENGTH(i.DNI) - 1), '-', RIGHT(i.DNI, 1)) AS rut,
  AT.nombreTrabajador AS NombreTrabajador,

  CASE i.obra_id
    WHEN 14168 THEN 'PG CD'
    WHEN 14184 THEN 'BAT LO BOZA'
    WHEN 14186 THEN 'UL CD'
    WHEN 14367 THEN 'PG VAS'
    WHEN 14368 THEN 'BAT CASABLANCA'
    WHEN 14369 THEN 'UL VAS'
    WHEN 14370 THEN 'PG BMP'
    WHEN 14818 THEN 'NOVICIADO'
    WHEN 16256 THEN 'PANAMERICANA'
    ELSE 'OTRO'
  END AS recinto,

  at.cuenta_area AS Cuenta,
  nc.cargo_normalizado AS Cargo,
  DATE_FORMAT(i.fecha_inasistencia, '%d/%m/%Y ') AS FECHA,

  CASE
    WHEN i.motivo = '-'        THEN 'Ausente'
    WHEN UPPER(i.motivo) = 'V' THEN 'Vacaciones'
    WHEN UPPER(i.motivo) = 'L' THEN 'Licencia'
    WHEN UPPER(i.motivo) = 'P' THEN 'Permiso'
    WHEN i.motivo IS NULL      THEN 'Sin registro'
    ELSE i.motivo
  END AS motivo,

  -- Campos técnicos para filtro/orden
  i.obra_id AS recinto_id,
  i.fecha_inasistencia AS fecha_real

FROM inasistencias i
JOIN asignacion_turnos at
  ON i.uid_inasistencia = at.uid_rut_dia_obra
JOIN nomina_colaborador nc
  ON i.DNI = nc.DNI
WHERE
  AT.tipoTurno IS NOT NULL
  AND i.obra_id = nc.obra_id
"""

SQL_HORAS_TRABAJADAS_BASE = """
SELECT DISTINCT
  CONCAT_WS(' ', a.nombre, a.apellido_paterno, a.apellido_materno)  AS NombreTrabajador,
  a.rut_trabajador                                                  AS dni,
  a.nombre_recinto                                                  AS recinto,
  at.diaTurno                                                       AS DiaTurno,
  a.entrada                                                         AS entrada,
  COALESCE(a.salida, '1999-01-01 00:00:00')                         AS salida,
  COALESCE(he.horas_total, 0)                                       AS HorasExtras,
  ROUND(
    GREATEST(
      TIMESTAMPDIFF(MINUTE, a.entrada, COALESCE(a.salida, a.salida_turno)),
      0
    ) / 60.0, 2
  )                                                                 AS HorasTrabajadas,
  ROUND(
    COALESCE(he.horas_total, 0) +
    GREATEST(TIMESTAMPDIFF(MINUTE, a.entrada, COALESCE(a.salida, a.salida_turno)), 0) / 60.0
  , 2)                                                              AS HorasTotal,
  a.entrada_turno                                                   AS entradaProgramada,
  a.salida_turno                                                    AS SalidaProgramada,
  a.cargo_resumido                                                  AS Cargo,
  at.tipoTurno                                                      AS tipo_turno,
  at.cuenta_area                                                    AS cuenta_area
FROM asistencia a
JOIN horas_extras_diario he
  ON a.rut_fecha_recinto = he.dni_fecha_recinto
LEFT JOIN asignacion_turnos at
  ON a.rut_fecha_recinto = at.uid_rut_dia_obra
"""


# ================= Rutas de navegación (HTML) =================

@bp.get("/")
@login_required
def index():
    return redirect(url_for("dashboard.dashboard"))


# =================== DASHBOARD ===================

@bp.get("/dashboard")
@login_required
def dashboard():
    # Filtros
    f_desde = (request.args.get("desde") or date.today().isoformat())
    f_hasta = (request.args.get("hasta") or date.today().isoformat())
    obra_id = (request.args.get("obra_id") or "").strip()
    cargo = (request.args.get("cargo") or "").strip()
    cuenta_area = (request.args.get("cuenta_area") or "").strip()

    # Acceso por recintos (obra_id)
    allowed = None if _user_is_super() else set(get_allowed_recintos_obra_ids(current_user.id))
    extra_asist, p_asist = _sql_in_clause_text("a.id_recinto", allowed)
    extra_inas,  p_inas  = _sql_in_clause_text("i.obra_id",    allowed)
    extra_combo, p_combo = _sql_in_clause_text("value",        allowed)

    # Permisos por cuentas para combo (pares recinto-cuenta)
    per_ctas = _allowed_cuentas(current_user.id, allowed)
    extra_cta_combo, p_cta_combo = _clause_cuentas("a.id_recinto", "a.cuenta_area", per_ctas)

    where_asist = ["DATE(a.fecha_base) BETWEEN :desde AND :hasta" + extra_asist]
    where_inas  = ["DATE(i.fecha_inasistencia) BETWEEN :desde AND :hasta" + extra_inas]
    params = {"desde": f_desde, "hasta": f_hasta, **p_asist, **p_inas}

    if obra_id:
        where_asist.append("a.id_recinto = :obra_id")
        where_inas.append("i.obra_id = :obra_id")
        params["obra_id"] = obra_id
    if cargo:
        params["cargo"] = cargo
        where_asist.append("a.cargo_resumido = :cargo")
    if cuenta_area:
        params["cuenta_area"] = cuenta_area
        where_asist.append("a.cuenta_area = :cuenta_area")

    wa, wi = " AND ".join(where_asist), " AND ".join(where_inas)

    # KPIs
    sql_kpis = text(f"""
        SELECT
          COALESCE(SUM(t.asistencias),0)   AS asistencia,
          COALESCE(SUM(t.inasistencias),0) AS inasistencia,
          COALESCE(SUM(t.asistencias)+SUM(t.inasistencias),0) AS dotacion,
          ROUND(COALESCE(SUM(t.asistencias),0) / NULLIF(COALESCE(SUM(t.asistencias),0)+COALESCE(SUM(t.inasistencias),0),0) * 100, 2) AS pct_presentismo,
          ROUND(COALESCE(SUM(t.inasistencias),0) / NULLIF(COALESCE(SUM(t.asistencias),0)+COALESCE(SUM(t.inasistencias),0),0) * 100, 2) AS pct_ausencia
        FROM (
          SELECT COUNT(*) AS asistencias, 0 AS inasistencias FROM asistencia a WHERE {wa} AND a.entrada IS NOT NULL
          UNION ALL
          SELECT 0 AS asistencias, COUNT(*) AS inasistencias FROM inasistencias i WHERE {wi}
        ) t;
    """)

    # Combos (obras, cargos y cuentas)
    sql_obras = text(f"""
        SELECT value, label
        FROM (
          SELECT DISTINCT a.id_recinto AS value,
            CASE a.id_recinto
              WHEN 14168 THEN 'PG CD' WHEN 14184 THEN 'BAT LO BOZA' WHEN 14186 THEN 'UL CD'
              WHEN 14367 THEN 'PG VAS' WHEN 14368 THEN 'BAT CASABLANCA' WHEN 14369 THEN 'UL VAS'
              WHEN 14370 THEN 'PG BMP' WHEN 14818 THEN 'NOVICIADO'     WHEN 16256 THEN 'PANAMERICANA'
              ELSE CONCAT('OBRA ', a.id_recinto) END AS label
          FROM asistencia a
          UNION
          SELECT DISTINCT i.obra_id AS value,
            CASE i.obra_id
              WHEN 14168 THEN 'PG CD' WHEN 14184 THEN 'BAT LO BOZA' WHEN 14186 THEN 'UL CD'
              WHEN 14367 THEN 'PG VAS' WHEN 14368 THEN 'BAT CASABLANCA' WHEN 14369 THEN 'UL VAS'
              WHEN 14370 THEN 'PG BMP' WHEN 14818 THEN 'NOVICIADO'     WHEN 16256 THEN 'PANAMERICANA'
              ELSE CONCAT('OBRA ', i.obra_id) END AS label
          FROM inasistencias i
        ) x
        WHERE 1=1 {extra_combo}
        ORDER BY label;
    """)

    sql_cargos = text("""
        SELECT DISTINCT a.cargo_resumido AS value
        FROM asistencia a
        WHERE a.cargo_resumido IS NOT NULL AND a.cargo_resumido <> ''
        ORDER BY value;
    """)

    # Combo cuentas filtrado por recintos *y* por cuentas asignadas
    sql_cuentas = text(f"""
        SELECT DISTINCT a.cuenta_area AS value
        FROM asistencia a
        WHERE a.cuenta_area IS NOT NULL
          AND a.cuenta_area <> ''
          {extra_asist}
          {extra_cta_combo}
        ORDER BY value;
    """)

    # Otras consultas de tablero (motivos, inas por recinto, resumen por recinto)
    sql_motivos = text(f"""
        SELECT motivo, cantidad,
               ROUND(cantidad / NULLIF(SUM(cantidad) OVER(), 0) * 100, 1) AS pct
        FROM (
            SELECT 
                CASE
                    WHEN i.motivo = '-'        THEN 'Ausentes'
                    WHEN UPPER(i.motivo) = 'V' THEN 'Vacaciones'
                    WHEN UPPER(i.motivo) = 'L' THEN 'Licencias'
                    WHEN UPPER(i.motivo) = 'P' THEN 'Permisos'
                    WHEN UPPER(i.motivo) = 'C' THEN 'Compensado'
                    WHEN i.motivo IS NULL      THEN 'Sin registro'
                    ELSE 'Otros'
                END AS motivo,
                COUNT(*) AS cantidad
            FROM inasistencias i
            WHERE {wi}
            GROUP BY motivo
        ) t
        ORDER BY cantidad DESC;
    """)

    sql_inas_por_recinto = text(f"""
        SELECT recinto, cantidad,
               ROUND(cantidad / NULLIF(SUM(cantidad) OVER(), 0) * 100, 1) AS pct
        FROM (
            SELECT 
                CASE i.obra_id
                    WHEN 14168 THEN 'PG CD' WHEN 14184 THEN 'BAT LO BOZA' WHEN 14186 THEN 'UL CD'
                    WHEN 14367 THEN 'PG VAS' WHEN 14368 THEN 'BAT CASABLANCA' WHEN 14369 THEN 'UL VAS'
                    WHEN 14370 THEN 'PG BMP' WHEN 14818 THEN 'NOVICIADO'     WHEN 16256 THEN 'PANAMERICANA'
                    ELSE CONCAT('OBRA ', i.obra_id)
                END AS recinto,
                COUNT(*) AS cantidad
            FROM inasistencias i
            WHERE {wi}
            GROUP BY i.obra_id
        ) t
        ORDER BY cantidad DESC;
    """)

    sql_resumen_recinto = text(f"""
        WITH
        asist AS (
            SELECT a.id_recinto AS recinto_id, COUNT(*) AS asist
            FROM asistencia a
            WHERE {wa} AND a.entrada IS NOT NULL
            GROUP BY a.id_recinto
        ),
        inas AS (
            SELECT i.obra_id AS recinto_id, COUNT(*) AS inasist
            FROM inasistencias i
            WHERE {wi}
            GROUP BY i.obra_id
        ),
        lic AS (
            SELECT i.obra_id AS recinto_id, COUNT(*) AS licencias
            FROM inasistencias i
            WHERE {wi} AND UPPER(i.motivo) = 'L'
            GROUP BY i.obra_id
        ),
        per AS (
            SELECT i.obra_id AS recinto_id, COUNT(*) AS permisos
            FROM inasistencias i
            WHERE {wi} AND UPPER(i.motivo) = 'P'
            GROUP BY i.obra_id
        ),
        vac AS (
            SELECT i.obra_id AS recinto_id, COUNT(*) AS vacaciones
            FROM inasistencias i
            WHERE {wi} AND UPPER(i.motivo) = 'V'
            GROUP BY i.obra_id
        )
        SELECT
            CASE r_id
                WHEN 14168 THEN 'PG CD' WHEN 14184 THEN 'BAT LO BOZA' WHEN 14186 THEN 'UL CD'
                WHEN 14367 THEN 'PG VAS' WHEN 14368 THEN 'BAT CASABLANCA' WHEN 14369 THEN 'UL VAS'
                WHEN 14370 THEN 'PG BMP' WHEN 14818 THEN 'NOVICIADO'     WHEN 16256 THEN 'PANAMERICANA'
                ELSE CONCAT('OBRA ', r_id)
            END AS recinto,
            COALESCE(a.asist, 0)     AS asist,
            COALESCE(i.inasist, 0)   AS inasist,
            COALESCE(a.asist, 0) + COALESCE(i.inasist, 0) AS dotacion,
            ROUND(COALESCE(a.asist,0) / NULLIF(COALESCE(a.asist,0) + COALESCE(i.inasist,0),0) * 100, 1) AS pct_pres,
            ROUND(COALESCE(i.inasist,0) / NULLIF(COALESCE(a.asist,0) + COALESCE(i.inasist,0),0) * 100, 1) AS pct_aus,
            COALESCE(l.licencias, 0) AS licencias,
            COALESCE(p.permisos, 0)  AS permisos,
            COALESCE(v.vacaciones, 0) AS vacaciones
        FROM (
            SELECT DISTINCT id_recinto AS r_id FROM asistencia a WHERE {wa}
            UNION
            SELECT DISTINCT obra_id   AS r_id FROM inasistencias i WHERE {wi}
        ) r
        LEFT JOIN asist a ON a.recinto_id = r.r_id
        LEFT JOIN inas  i ON i.recinto_id = r.r_id
        LEFT JOIN lic   l ON l.recinto_id = r.r_id
        LEFT JOIN per   p ON p.recinto_id = r.r_id
        LEFT JOIN vac   v ON v.recinto_id = r.r_id
        ORDER BY recinto ASC;
    """)

    with db.engine.begin() as conn:
        k = conn.execute(sql_kpis, params).mappings().first() or {}
        obras = conn.execute(sql_obras, dict(p_combo)).mappings().all()
        cargos = conn.execute(sql_cargos).mappings().all()
        cuentas = conn.execute(sql_cuentas, {**p_asist, **p_cta_combo}).mappings().all()
        motivos = conn.execute(sql_motivos, params).mappings().all()
        inas_por_recinto = conn.execute(sql_inas_por_recinto, params).mappings().all()
        resumen_recinto = conn.execute(sql_resumen_recinto, params).mappings().all()

    kpis = {
        "asistencia": int(k.get("asistencia", 0)),
        "inasistencia": int(k.get("inasistencia", 0)),
        "dotacion": int(k.get("dotacion", 0)),
        "pct_presentismo": float(k.get("pct_presentismo") or 0.0),
        "pct_ausencia": float(k.get("pct_ausencia") or 0.0),
    }
    filtros = {"desde": f_desde, "hasta": f_hasta, "obra_id": obra_id, "cargo": cargo, "cuenta_area": cuenta_area}
    opciones = {"obras": obras, "cargos": cargos, "cuentas": cuentas}

    return render_template(
        "dashboard/index.html",
        kpis=kpis,
        tabla=[],
        filtros=filtros,
        opciones=opciones,
        motivos=motivos,
        inas_por_recinto=inas_por_recinto,
        resumen_recinto=resumen_recinto
    )


# =================== REPORTES ===================

@bp.get("/reporte/horas-trabajadas")
@login_required
def reporte_horas_trabajadas():
    start = (request.args.get("start") or "").strip()
    end   = (request.args.get("end") or "").strip()
    if not start or not end:
        end_date   = _date.today()
        start_date = end_date - _timedelta(days=6)
        start = start_date.isoformat()
        end   = end_date.isoformat()
    return render_template("dashboard/reporte_horas_trabajadas.html", start=start, end=end)


@bp.get("/api/horas-trabajadas")
@login_required
def api_horas_trabajadas():
    start = (request.args.get("start") or "").strip()
    end   = (request.args.get("end") or "").strip()
    if not start or not end:
        return jsonify({"error": "Parámetros 'start' y 'end' son obligatorios (YYYY-MM-DD)."}), 400

    try: page = max(int(request.args.get("page", 1)), 1)
    except: page = 1
    try: per_page = int(request.args.get("per_page", 30))
    except: per_page = 30
    per_page = max(1, min(per_page, 200))
    offset = (page - 1) * per_page

    allowed = _allowed_recinto_ids()
    extra_asist, p_asist = _sql_in_clause_text("a.id_recinto", allowed)
    # filtro por cuentas
    per_ctas = _allowed_cuentas(current_user.id, allowed)
    extra_cta, p_cta = _clause_cuentas("a.id_recinto", "a.cuenta_area", per_ctas)

    WHERE = f"""
      WHERE DATE(a.fecha_base) BETWEEN :start AND :end
        AND at.tipoTurno IS NOT NULL
        {extra_asist}
        {extra_cta}
    """
    params = {"start": start, "end": end, **p_asist, **p_cta}

    sql_count = text(f"SELECT COUNT(*) FROM ({SQL_HORAS_TRABAJADAS_BASE} {WHERE}) AS q")
    total = db.session.execute(sql_count, params).scalar() or 0

    sql_page = text(f"""
      SELECT
        t.NombreTrabajador, t.dni, t.recinto,
        DATE_FORMAT(t.DiaTurno,'%d/%m/%Y')                    AS DiaTurno,
        DATE_FORMAT(t.entrada, '%d/%m/%Y %H:%i:%s')           AS entrada,
        DATE_FORMAT(t.salida,  '%d/%m/%Y %H:%i:%s')           AS salida,
        t.HorasExtras, t.HorasTrabajadas, t.HorasTotal,
        DATE_FORMAT(t.entradaProgramada, '%d/%m/%Y %H:%i:%s') AS entradaProgramada,
        DATE_FORMAT(t.SalidaProgramada,  '%d/%m/%Y %H:%i:%s') AS SalidaProgramada,
        t.Cargo, t.tipo_turno, t.cuenta_area
      FROM (
        {SQL_HORAS_TRABAJADAS_BASE}
        {WHERE}
      ) t
      ORDER BY t.recinto, t.dni, t.entrada
      LIMIT :limit OFFSET :offset
    """)
    rows = db.session.execute(sql_page, {**params, "limit": per_page, "offset": offset}).mappings().all()
    items = [dict(r) for r in rows]

    pages = (total // per_page) + (1 if total % per_page else 0)
    return jsonify({
        "items": items, "page": page, "per_page": per_page,
        "total": total, "pages": pages,
        "has_prev": page > 1, "has_next": page < pages if pages else False
    })


@bp.get("/reporte/horas-trabajadas/export")
@login_required
def export_horas_trabajadas():
    start = (request.args.get("start") or "").strip()
    end   = (request.args.get("end") or "").strip()
    if not start or not end:
        return "Parámetros 'start' y 'end' son obligatorios (YYYY-MM-DD).", 400

    allowed = _allowed_recinto_ids()
    extra_asist, p_asist = _sql_in_clause_text("a.id_recinto", allowed)
    per_ctas = _allowed_cuentas(current_user.id, allowed)
    extra_cta, p_cta = _clause_cuentas("a.id_recinto", "a.cuenta_area", per_ctas)

    WHERE = f"""
      WHERE DATE(a.fecha_base) BETWEEN :start AND :end
        AND at.tipoTurno IS NOT NULL
        {extra_asist}
        {extra_cta}
    """

    sql_all = text(f"""
      SELECT
        t.NombreTrabajador, t.dni, t.recinto,
        DATE_FORMAT(t.DiaTurno,'%d/%m/%Y')                    AS DiaTurno,
        DATE_FORMAT(t.entrada, '%d/%m/%Y %H:%i:%s')           AS entrada,
        DATE_FORMAT(t.salida,  '%d/%m/%Y %H:%i:%s')           AS salida,
        t.HorasExtras, t.HorasTrabajadas, t.HorasTotal,
        DATE_FORMAT(t.entradaProgramada, '%d/%m/%Y %H:%i:%s') AS entradaProgramada,
        DATE_FORMAT(t.SalidaProgramada,  '%d/%m/%Y %H:%i:%s') AS SalidaProgramada,
        t.Cargo, t.tipo_turno, t.cuenta_area
      FROM (
        {SQL_HORAS_TRABAJADAS_BASE}
        {WHERE}
      ) t
      ORDER BY t.recinto, t.dni, t.entrada
    """)
    rows = db.session.execute(sql_all, {"start": start, "end": end, **p_asist, **p_cta}).mappings().all()

    df = pd.DataFrame([dict(r) for r in rows])
    cols = [
        "NombreTrabajador","dni","recinto","DiaTurno","entrada","salida",
        "HorasExtras","HorasTrabajadas","HorasTotal",
        "entradaProgramada","SalidaProgramada","Cargo","tipo_turno","cuenta_area"
    ]
    if df.empty:
        df = pd.DataFrame(columns=cols)
    else:
        df = df.reindex(columns=cols)

    buf = BytesIO()
    fname = f"horas_trabajadas_{start}_a_{end}"

    try:
        import xlsxwriter; engine = "xlsxwriter"
    except ImportError:
        try:
            import openpyxl; engine = "openpyxl"
        except ImportError:
            engine = None

    if engine:
        with pd.ExcelWriter(buf, engine=engine) as writer:
            df.to_excel(writer, index=False, sheet_name="Horas")
        buf.seek(0)
        return send_file(buf, as_attachment=True,
                         download_name=f"{fname}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    df.to_csv(buf, index=False, encoding="utf-8-sig"); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{fname}.csv", mimetype="text/csv")


# ---- Inasistencias (HTML, API y export) ----

@bp.get("/reporte/inasistencias")
@login_required
def reporte_inasistencias():
    return render_template("dashboard/reporte_inasistencias.html", page_title="Inasistencias")


@bp.get("/api/inasistencias")
@login_required
def api_inasistencias():
    start = (request.args.get("start") or "").strip()
    end   = (request.args.get("end") or "").strip()
    if not start or not end:
        return jsonify({"error": "Parámetros 'start' y 'end' son obligatorios (YYYY-MM-DD)."}), 400

    try: page = max(int(request.args.get("page", 1)), 1)
    except: page = 1
    try: per_page = int(request.args.get("per_page", 30))
    except: per_page = 30
    per_page = max(1, min(per_page, 200))
    offset = (page - 1) * per_page

    allowed = _allowed_recinto_ids()
    extra_sql, extra_params = _sql_in_clause_text("t.recinto_id", allowed)

    # filtro por cuentas dentro del subquery base (i.obra_id + at.cuenta_area)
    per_ctas = _allowed_cuentas(current_user.id, allowed)
    extra_cta, p_cta = _clause_cuentas("i.obra_id", "at.cuenta_area", per_ctas)

    WHERE = f"WHERE t.fecha_real BETWEEN :start AND :end{extra_sql}"
    params = {"start": start, "end": end, **extra_params, **p_cta}

    sql_count = text(f"SELECT COUNT(*) FROM ( {SQL_INASISTENCIAS_BASE} {extra_cta} ) t {WHERE}")
    total = db.session.execute(sql_count, params).scalar() or 0

    sql_page = text(f"""
        SELECT * FROM ( {SQL_INASISTENCIAS_BASE} {extra_cta} ) t
        {WHERE}
        ORDER BY t.fecha_real DESC, t.recinto, t.rut
        LIMIT :limit OFFSET :offset
    """)
    rows = db.session.execute(sql_page, {**params, "limit": per_page, "offset": offset}).mappings().all()

    items = []
    for r in rows:
        d = dict(r); d.pop("fecha_real", None); d.pop("recinto_id", None)
        items.append(d)

    return jsonify({
        "items": items,
        "page": page,
        "per_page": per_page,
        "total": total,
        "pages": (total // per_page) + (1 if total % per_page else 0)
    })


@bp.get("/reporte/inasistencias/export")
@login_required
def export_inasistencias():
    start = (request.args.get("start") or "").strip()
    end   = (request.args.get("end") or "").strip()
    if not start or not end:
        return "Parámetros 'start' y 'end' son obligatorios (YYYY-MM-DD).", 400

    allowed = _allowed_recinto_ids()
    extra_sql, extra_params = _sql_in_clause_text("t.recinto_id", allowed)

    per_ctas = _allowed_cuentas(current_user.id, allowed)
    extra_cta, p_cta = _clause_cuentas("i.obra_id", "at.cuenta_area", per_ctas)

    sql = text(f"""
        SELECT * FROM ( {SQL_INASISTENCIAS_BASE} {extra_cta} ) t
        WHERE t.fecha_real BETWEEN :start AND :end{extra_sql}
        ORDER BY t.fecha_real DESC, t.recinto, t.rut
    """)
    rows = db.session.execute(sql, {"start": start, "end": end, **extra_params, **p_cta}).mappings().all()

    df = pd.DataFrame([dict(r) for r in rows]).drop(columns=["fecha_real", "recinto_id"], errors="ignore")
    buf = BytesIO()
    fname = f"inasistencias_{start}_a_{end}"

    try:
        import xlsxwriter; engine = "xlsxwriter"
    except Exception:
        try: import openpyxl; engine = "openpyxl"
        except Exception: engine = None

    if engine:
        with pd.ExcelWriter(buf, engine=engine) as w:
            df.to_excel(w, index=False, sheet_name="Inasistencias")
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f"{fname}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    df.to_csv(buf, index=False, encoding="utf-8-sig"); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{fname}.csv", mimetype="text/csv")


# =================== HORAS EXTRA ===================

@bp.get("/reporte/horas-extras")
@login_required
def reporte_horas_extras():
    end_date   = _date.today()
    start_date = end_date - _timedelta(days=6)
    return render_template("dashboard/reporte_horas_extras.html",
                           start=start_date.isoformat(),
                           end=end_date.isoformat())


@bp.get("/api/horas-extras")
@login_required
def api_horas_extras():
    start = (request.args.get("start") or "").strip()
    end   = (request.args.get("end") or "").strip()
    if not start or not end:
        return jsonify({"error": "Parámetros 'start' y 'end' son obligatorios (YYYY-MM-DD)."}), 400

    try: page = max(int(request.args.get("page", 1)), 1)
    except: page = 1
    try: per_page = int(request.args.get("per_page", 30))
    except: per_page = 30
    per_page = max(1, min(per_page, 200))
    offset = (page - 1) * per_page

    # --- permisos por recintos
    allowed = _allowed_recinto_ids()
    extra_asist, p_asist = _sql_in_clause_text("a.id_recinto", allowed)

    # --- NUEVO: permisos por cuentas (pares recinto-cuenta)
    per_ctas = _allowed_cuentas(current_user.id, allowed)
    extra_cta, p_cta = _clause_cuentas("a.id_recinto", "a.cuenta_area", per_ctas)

    WHERE = f"""
      WHERE DATE(he.fecha) BETWEEN :start AND :end
        {extra_asist}
        {extra_cta}
    """
    params = {"start": start, "end": end, **p_asist, **p_cta}

    sql_count = text(f"""
      SELECT COUNT(*) FROM (
        SELECT he.dni_fecha_recinto
        FROM horas_extras_diario he
        LEFT JOIN asistencia a
          ON a.rut_fecha_recinto = he.dni_fecha_recinto
        {WHERE}
        GROUP BY he.dni_fecha_recinto, he.fecha
      ) q
    """)
    total = db.session.execute(sql_count, params).scalar() or 0

    sql_page = text(f"""
      SELECT
        CONCAT_WS(' ', a.nombre, a.apellido_paterno, a.apellido_materno)  AS NombreTrabajador,
        a.rut_trabajador                                                  AS dni,
        a.nombre_recinto                                                  AS recinto,
        DATE_FORMAT(he.fecha, '%d/%m/%Y')                                 AS fecha,
        a.cargo_resumido                                                  AS cargo,
        a.cuenta_area                                                     AS cuenta_area,
        ROUND(SUM(he.horas_total), 2)                                     AS horas_extras
      FROM horas_extras_diario he
      LEFT JOIN asistencia a
        ON a.rut_fecha_recinto = he.dni_fecha_recinto
      {WHERE}
      GROUP BY a.rut_trabajador, a.nombre, a.apellido_paterno, a.apellido_materno,
               a.nombre_recinto, a.cargo_resumido, a.cuenta_area, he.fecha
      ORDER BY he.fecha DESC, recinto, dni
      LIMIT :limit OFFSET :offset
    """)
    rows = db.session.execute(sql_page, {**params, "limit": per_page, "offset": offset}).mappings().all()

    return jsonify({
        "items": [dict(r) for r in rows],
        "page": page, "per_page": per_page, "total": total,
        "pages": (total // per_page) + (1 if total % per_page else 0)
    })


@bp.get("/reporte/horas-extras/export")
@login_required
def export_horas_extras():
    start = (request.args.get("start") or "").strip()
    end   = (request.args.get("end") or "").strip()
    if not start or not end:
        return "Parámetros 'start' y 'end' son obligatorios (YYYY-MM-DD).", 400

    # --- permisos por recintos
    allowed = _allowed_recinto_ids()
    extra_asist, p_asist = _sql_in_clause_text("a.id_recinto", allowed)

    # --- NUEVO: permisos por cuentas (pares recinto-cuenta)
    per_ctas = _allowed_cuentas(current_user.id, allowed)
    extra_cta, p_cta = _clause_cuentas("a.id_recinto", "a.cuenta_area", per_ctas)

    WHERE = f"""
      WHERE DATE(he.fecha) BETWEEN :start AND :end
        {extra_asist}
        {extra_cta}
    """

    sql_all = text(f"""
      SELECT
        CONCAT_WS(' ', a.nombre, a.apellido_paterno, a.apellido_materno)  AS NombreTrabajador,
        a.rut_trabajador                                                  AS dni,
        a.nombre_recinto                                                  AS recinto,
        DATE_FORMAT(he.fecha, '%d/%m/%Y')                                 AS fecha,
        a.cargo_resumido                                                  AS cargo,
        a.cuenta_area                                                     AS cuenta_area,
        ROUND(SUM(he.horas_total), 2)                                     AS horas_extras
      FROM horas_extras_diario he
      LEFT JOIN asistencia a
        ON a.rut_fecha_recinto = he.dni_fecha_recinto
      {WHERE}
      GROUP BY a.rut_trabajador, a.nombre, a.apellido_paterno, a.apellido_materno,
               a.nombre_recinto, a.cargo_resumido, a.cuenta_area, he.fecha
      ORDER BY he.fecha DESC, recinto, dni
    """)
    rows = db.session.execute(sql_all, {"start": start, "end": end, **p_asist, **p_cta}).mappings().all()

    df = pd.DataFrame([dict(r) for r in rows])
    cols = ["fecha","dni","NombreTrabajador","recinto","cargo","cuenta_area","horas_extras"]
    if df.empty:
        df = pd.DataFrame(columns=cols)
    else:
        df = df.reindex(columns=cols)

    buf = BytesIO()
    fname = f"horas_extras_{start}_a_{end}"
    try:
        import xlsxwriter; engine = "xlsxwriter"
    except Exception:
        try:
            import openpyxl; engine = "openpyxl"
        except Exception:
            engine = None

    if engine:
        with pd.ExcelWriter(buf, engine=engine) as writer:
            df.to_excel(writer, index=False, sheet_name="Horas extra")
        buf.seek(0)
        return send_file(buf, as_attachment=True,
                         download_name=f"{fname}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    df.to_csv(buf, index=False, encoding="utf-8-sig"); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{fname}.csv", mimetype="text/csv")


# =================== Presentismo ===================

@bp.get("/presentismo")
@login_required
def presentismo():
    def _pct2(p, a):
        p = int(p or 0); a = int(a or 0); t = p + a
        return round(100.0 * p / t, 1) if t else 0.0

    f_desde = (request.args.get("desde") or "").strip()
    f_hasta = (request.args.get("hasta") or "").strip()
    rid     = request.args.get("rid", type=int)

    today = date.today()
    if not f_desde:
        f_desde = today.replace(day=1).isoformat()
    if not f_hasta:
        f_hasta = today.isoformat()

    weekday    = today.weekday()
    week_start = (today - timedelta(days=weekday)).isoformat()
    week_end   = (today + timedelta(days=(6 - weekday))).isoformat()

    allowed = None if _user_is_super() else set(get_allowed_recintos_obra_ids(current_user.id))
    if rid:
        enforce_rid_allowed(rid)

    def _expand_in(colname, allowed_set):
        if allowed_set is None:
            return "", {}
        if not allowed_set:
            return " AND 1=0 ", {}
        return f" AND {colname} IN :rids ", {"rids": tuple(allowed_set)}

    extraA, pA = _expand_in("a.id_recinto", allowed if not rid else None)
    extraI, pI = _expand_in("i.obra_id",    allowed if not rid else None)

    params_rango = {"desde": f_desde, "hasta": f_hasta, "rid": rid, **pA, **pI}
    params_mes   = {"mes_ini": f_desde[:7]+"-01", "mes_fin": f_hasta, "rid": rid, **pA, **pI}
    params_sem   = {"sem_ini": week_start, "sem_fin": week_end, "rid": rid, **pA, **pI}
    params_year  = {"y": int(f_hasta[:4]), "rid": rid, **pA, **pI}

    sql_mes = text(f"""
        SELECT SUM(presentes) AS presentes, SUM(ausentes) AS ausentes
        FROM (
          SELECT COUNT(*) AS presentes, 0 AS ausentes
          FROM asistencia a
          WHERE DATE(a.fecha_base) BETWEEN :mes_ini AND :mes_fin
            AND a.entrada IS NOT NULL
            AND (:rid IS NULL OR a.id_recinto = :rid)
            {extraA}
          UNION ALL
          SELECT 0 AS presentes, COUNT(*) AS ausentes
          FROM inasistencias i
          WHERE DATE(i.fecha_inasistencia) BETWEEN :mes_ini AND :mes_fin
            AND (:rid IS NULL OR i.obra_id = :rid)
            {extraI}
        ) t
    """)

    sql_sem = text(f"""
        SELECT SUM(presentes) AS presentes, SUM(ausentes) AS ausentes
        FROM (
          SELECT COUNT(*) AS presentes, 0 AS ausentes
          FROM asistencia a
          WHERE DATE(a.fecha_base) BETWEEN :sem_ini AND :sem_fin
            AND a.entrada IS NOT NULL
            AND (:rid IS NULL OR a.id_recinto = :rid)
            {extraA}
          UNION ALL
          SELECT 0 AS presentes, COUNT(*) AS ausentes
          FROM inasistencias i
          WHERE DATE(i.fecha_inasistencia) BETWEEN :sem_ini AND :sem_fin
            AND (:rid IS NULL OR i.obra_id = :rid)
            {extraI}
        ) t
    """)

    sql_recinto_uno = text("""
        SELECT
          :rid AS rid,
          CASE :rid
            WHEN 14168 THEN 'PG CD' WHEN 14184 THEN 'BAT LO BOZA' WHEN 14186 THEN 'UL CD'
            WHEN 14367 THEN 'PG VAS' WHEN 14368 THEN 'BAT CASABLANCA' WHEN 14369 THEN 'UL VAS'
            WHEN 14370 THEN 'PG BMP' WHEN 14818 THEN 'NOVICIADO'     WHEN 16256 THEN 'PANAMERICANA'
            ELSE CONCAT('OBRA ', :rid)
          END AS recinto,
          (SELECT COUNT(*) FROM asistencia a
            WHERE DATE(a.fecha_base) BETWEEN :desde AND :hasta
              AND a.entrada IS NOT NULL
              AND a.id_recinto = :rid) AS presentes,
          (SELECT COUNT(*) FROM inasistencias i
            WHERE DATE(i.fecha_inasistencia) BETWEEN :desde AND :hasta
              AND i.obra_id = :rid) AS ausentes
    """)

    sql_recintos_ranking = text(f"""
        WITH pres AS (
          SELECT a.id_recinto AS rid, COUNT(*) AS presentes
          FROM asistencia a
          WHERE DATE(a.fecha_base) BETWEEN :desde AND :hasta
            AND a.entrada IS NOT NULL
            {extraA}
          GROUP BY a.id_recinto
        ),
        aus AS (
          SELECT i.obra_id AS rid, COUNT(*) AS ausentes
          FROM inasistencias i
          WHERE DATE(i.fecha_inasistencia) BETWEEN :desde AND :hasta
            {extraI}
          GROUP BY i.obra_id
        ),
        rids AS ( SELECT rid FROM pres UNION SELECT rid FROM aus )
        SELECT
          r.rid AS rid,
          CASE r.rid
            WHEN 14168 THEN 'PG CD' WHEN 14184 THEN 'BAT LO BOZA' WHEN 14186 THEN 'UL CD'
            WHEN 14367 THEN 'PG VAS' WHEN 14368 THEN 'BAT CASABLANCA' WHEN 14369 THEN 'UL VAS'
            WHEN 14370 THEN 'PG BMP' WHEN 14818 THEN 'NOVICIADO'     WHEN 16256 THEN 'PANAMERICANA'
            ELSE CONCAT('OBRA ', r.rid)
          END AS recinto,
          COALESCE(p.presentes,0) AS presentes,
          COALESCE(a.ausentes,0)  AS ausentes
        FROM rids r
        LEFT JOIN pres p ON p.rid = r.rid
        LEFT JOIN aus  a ON a.rid = r.rid
        ORDER BY (COALESCE(p.presentes,0)+COALESCE(a.ausentes,0)) DESC, recinto
    """)

    sql_meses = text(f"""
        WITH m_presentes AS (
          SELECT MONTH(a.fecha_base) AS m, COUNT(*) AS presentes
          FROM asistencia a
          WHERE YEAR(a.fecha_base)=:y AND a.entrada IS NOT NULL
            AND (:rid IS NULL OR a.id_recinto = :rid)
            {extraA}
          GROUP BY MONTH(a.fecha_base)
        ),
        m_ausentes AS (
          SELECT MONTH(i.fecha_inasistencia) AS m, COUNT(*) AS ausentes
          FROM inasistencias i
          WHERE YEAR(i.fecha_inasistencia)=:y
            AND (:rid IS NULL OR i.obra_id = :rid)
            {extraI}
          GROUP BY MONTH(i.fecha_inasistencia)
        ),
        cal AS (
          SELECT 1 AS mn UNION ALL SELECT 2 UNION ALL SELECT 3 UNION ALL SELECT 4 UNION ALL SELECT 5 UNION ALL SELECT 6
          UNION ALL SELECT 7 UNION ALL SELECT 8 UNION ALL SELECT 9 UNION ALL SELECT 10 UNION ALL SELECT 11 UNION ALL SELECT 12
        )
        SELECT
          LPAD(cal.mn,2,'0') AS mes_num,
          CASE cal.mn
            WHEN 1 THEN 'ene' WHEN 2 THEN 'feb' WHEN 3 THEN 'mar'
            WHEN 4 THEN 'abr' WHEN 5 THEN 'may' WHEN 6 THEN 'jun'
            WHEN 7 THEN 'jul' WHEN 8 THEN 'ago' WHEN 9 THEN 'sep'
            WHEN 10 THEN 'oct' WHEN 11 THEN 'nov' ELSE 'dic'
          END AS mes,
          COALESCE(p.presentes,0) AS presentes,
          COALESCE(a.ausentes,0)  AS ausentes
        FROM cal
        LEFT JOIN m_presentes p ON p.m = cal.mn
        LEFT JOIN m_ausentes a  ON a.m = cal.mn
        ORDER BY cal.mn
    """)

    with db.engine.begin() as conn:
        r_mes = conn.execute(sql_mes, params_mes).mappings().first() or {"presentes":0,"ausentes":0}
        r_sem = conn.execute(sql_sem, params_sem).mappings().first() or {"presentes":0,"ausentes":0}
        if rid:
            recs = conn.execute(sql_recinto_uno, params_rango).mappings().all()
        else:
            recs = conn.execute(sql_recintos_ranking, params_rango).mappings().all()
        meses = conn.execute(sql_meses, params_year).mappings().all()

    mes_p = int(r_mes["presentes"] or 0); mes_a = int(r_mes["ausentes"] or 0)
    sem_p = int(r_sem["presentes"] or 0); sem_a = int(r_sem["ausentes"] or 0)

    recintos_py = []
    for r in recs:
        presentes = int(r.get("presentes") or 0)
        ausentes  = int(r.get("ausentes")  or 0)
        rid_val   = int(r.get("rid") or 0)
        recinto_n = str(r.get("recinto") or "")
        total = presentes + ausentes
        pct_pres = round(presentes * 100.0 / total, 1) if total else 0.0
        pct_aus  = round(ausentes  * 100.0 / total, 1) if total else 0.0
        recintos_py.append({
            "rid": rid_val, "recinto": recinto_n,
            "presentes": presentes, "ausentes": ausentes,
            "total": total, "pct_pres": pct_pres, "pct_aus": pct_aus
        })

    meses_py = [{"mes": str(m["mes"]), "presentes": int(m["presentes"] or 0), "ausentes": int(m["ausentes"] or 0)} for m in meses]
    selected_name = recintos_py[0]["recinto"] if rid and recintos_py else None

    vm = {
        "filtros": {"desde": f_desde, "hasta": f_hasta, "sem_ini": week_start, "sem_fin": week_end},
        "rid": rid,
        "rid_name": selected_name,
        "donut_mes": {"presentes": mes_p, "ausentes": mes_a, "pct": (round(100.0 * mes_p / (mes_p+mes_a),1) if (mes_p+mes_a) else 0.0)},
        "donut_sem": {"presentes": sem_p, "ausentes": sem_a, "pct": (round(100.0 * sem_p / (sem_p+sem_a),1) if (sem_p+sem_a) else 0.0)},
        "recintos": recintos_py,
        "meses": meses_py,
    }
    return render_template("dashboard/presentismo.html", **vm)


# ========= Vistas simples =========

@bp.get("/reporte/movimientos")
@login_required
def reporte_movimientos():
    return render_template("dashboard/reporte_movimientos.html")

# =================== NOMINA ===================
@bp.get("/api/nomina/cuentas")
@login_required
def api_nomina_cuentas():
    """Devuelve las cuenta_area disponibles para el usuario (según permisos).
       Si no hay permisos explícitos, hace fallback a asistencia en el rango de fechas."""
    from datetime import date
    start = (request.args.get("start") or "").strip()
    end   = (request.args.get("end") or "").strip()

    # Rango por defecto: mes actual
    today = date.today()
    if not start:
        start = today.replace(day=1).isoformat()
    if not end:
        end = today.isoformat()

    allowed = _allowed_recinto_ids() or []
    per_ctas = _allowed_cuentas(current_user.id, allowed) or []

    cuentas_set = set()

    # Normalizador para cualquier forma de (_allowed_cuentas)
    def _extract_cta(item):
        # tupla/lista (rid, cta)
        if isinstance(item, (tuple, list)) and len(item) >= 2:
            return item[1]
        # dict {"recinto_id":..., "cuenta_area":...} o {"rid":..., "cta":...}
        if isinstance(item, dict):
            return item.get("cuenta_area") or item.get("cta") or item.get("account") or item.get("cuenta")
        # objeto con atributos
        c = getattr(item, "cuenta_area", None) or getattr(item, "cta", None) or getattr(item, "account", None)
        return c

    for it in per_ctas:
        cta = _extract_cta(it)
        if cta:
            cuentas_set.add(str(cta).strip())

    # Fallback: si no obtuvimos nada de permisos, miramos asistencia
    if not cuentas_set:
        extra_asist, params_asist = _sql_in_clause_text("a.id_recinto", allowed)
        sql = f"""
            SELECT DISTINCT a.cuenta_area
            FROM asistencia a
            WHERE a.cuenta_area IS NOT NULL AND a.cuenta_area <> ''
              AND DATE(a.fecha_base) BETWEEN :start AND :end
              {extra_asist}
            ORDER BY a.cuenta_area
        """
        rows = db.session.execute(text(sql), {"start": start, "end": end, **params_asist}).all()
        for (cta,) in rows:
            cuentas_set.add(cta)

    return jsonify(sorted(cuentas_set))





@bp.get("/nomina")
@login_required
def nomina():
    from datetime import date
    today = date.today()
    start = today.replace(day=1).isoformat()  # YYYY-MM-DD
    end = today.isoformat()                   # YYYY-MM-DD
    return render_template("dashboard/reporte_nomina.html", start=start, end=end)


@bp.get("/api/nomina")
@login_required
def api_nomina():
    start = (request.args.get("start") or "").strip()
    end   = (request.args.get("end") or "").strip()
    if not start or not end:
        return jsonify({"error": "Parámetros 'start' y 'end' son obligatorios (YYYY-MM-DD)."}), 400

    cuenta_area = (request.args.get("cuenta_area") or "").strip()

    # paginación
    page = max(1, request.args.get("page", type=int) or 1)
    per_page = request.args.get("per_page", type=int) or 100
    per_page = min(max(per_page, 10), 500)  # clamp 10..500
    offset = (page - 1) * per_page

    # permisos (igual que antes)
    allowed = _allowed_recinto_ids()
    extra_asist, p_asist = _sql_in_clause_text("a.id_recinto", allowed)
    extra_inas,  p_inas  = _sql_in_clause_text("i.obra_id", allowed)
    per_ctas = _allowed_cuentas(current_user.id, allowed)
    extra_cta_asist, p_cta_asist = _clause_cuentas("a.id_recinto", "a.cuenta_area", per_ctas)
    extra_cta_inas,  p_cta_inas  = _clause_cuentas("i.obra_id", "at.cuenta_area",  per_ctas)

    # CTE base (reutilizable para COUNT y PAGE)
    BASE = f"""
    WITH asist AS (
        SELECT
            a.rut_trabajador AS rut,
            CONCAT_WS(' ', a.nombre, a.apellido_paterno, a.apellido_materno) AS nombre_completo,
            a.nombre_recinto AS recinto,
            a.cuenta_area    AS cuenta_area,
            a.cargo_resumido AS cargo,
            COUNT(*) AS dias_asistidos
        FROM asistencia a
        WHERE DATE(a.fecha_base) BETWEEN :start AND :end
          AND a.entrada IS NOT NULL
          {extra_asist}
          {extra_cta_asist}
          AND (:cta = '' OR a.cuenta_area = :cta)
        GROUP BY a.rut_trabajador, a.nombre, a.apellido_paterno, a.apellido_materno,
                 a.nombre_recinto, a.cuenta_area, a.cargo_resumido
    ),
    inas AS (
        SELECT
            REPLACE(REPLACE(i.DNI,'.',''),'-','') AS dni_norm,
            at.cuenta_area AS cuenta_area,
            COUNT(*) AS dias_inasistentes
        FROM inasistencias i
        JOIN asignacion_turnos at
          ON i.uid_inasistencia = at.uid_rut_dia_obra
        WHERE DATE(i.fecha_inasistencia) BETWEEN :start AND :end
          {extra_inas}
          {extra_cta_inas}
          AND (:cta = '' OR at.cuenta_area = :cta)
        GROUP BY dni_norm, at.cuenta_area
    ),
    final AS (
        SELECT
            a.rut,
            a.nombre_completo AS nombre,
            a.recinto,
            a.cuenta_area,
            a.cargo,
            COALESCE(a.dias_asistidos,0)    AS dias_asistidos,
            COALESCE(i.dias_inasistentes,0) AS dias_inasistentes,
            (COALESCE(a.dias_asistidos,0) + COALESCE(i.dias_inasistentes,0)) AS total_dias,
            ROUND(COALESCE(a.dias_asistidos,0) / NULLIF((COALESCE(a.dias_asistidos,0) + COALESCE(i.dias_inasistentes,0)),0) * 100, 1) AS pct_asistencia
        FROM asist a
        LEFT JOIN inas i
          ON REPLACE(REPLACE(a.rut,'.',''),'-','') = i.dni_norm
         AND a.cuenta_area = i.cuenta_area
    )
    """

    # total
    COUNT_SQL = BASE + "SELECT COUNT(*) AS total FROM final;"
    params = {"start": start, "end": end, "cta": cuenta_area,
              **p_asist, **p_inas, **p_cta_asist, **p_cta_inas}
    total = db.session.execute(text(COUNT_SQL), params).scalar() or 0
    pages = (total + per_page - 1) // per_page if total else 0

    # página
    PAGE_SQL = BASE + """
    SELECT * FROM final
    ORDER BY recinto, cuenta_area, nombre
    LIMIT :limit OFFSET :offset;
    """
    rows = db.session.execute(text(PAGE_SQL), {**params, "limit": per_page, "offset": offset}).mappings().all()

    return jsonify({
        "items": [dict(r) for r in rows],
        "page": page, "per_page": per_page, "pages": pages,
        "total": total, "start": start, "end": end
    })


@bp.get("/nomina/export")
@login_required
def export_nomina():
    """
    Exporta la nómina consolidada a XLSX o CSV.
    """
    start = (request.args.get("start") or "").strip()
    end   = (request.args.get("end") or "").strip()
    if not start or not end:
        return "Parámetros 'start' y 'end' son obligatorios (YYYY-MM-DD).", 400

    cuenta_area = (request.args.get("cuenta_area") or "").strip()

    allowed = _allowed_recinto_ids()
    extra_asist, p_asist = _sql_in_clause_text("a.id_recinto", allowed)
    extra_inas,  p_inas  = _sql_in_clause_text("i.obra_id", allowed)
    per_ctas = _allowed_cuentas(current_user.id, allowed)
    extra_cta_asist, p_cta_asist = _clause_cuentas("a.id_recinto", "a.cuenta_area", per_ctas)
    extra_cta_inas,  p_cta_inas  = _clause_cuentas("i.obra_id", "at.cuenta_area", per_ctas)

    sql = f"""
    WITH asist AS (
        SELECT
            a.rut_trabajador AS rut,
            CONCAT_WS(' ', a.nombre, a.apellido_paterno, a.apellido_materno) AS nombre_completo,
            a.nombre_recinto AS recinto,
            a.cuenta_area    AS cuenta_area,
            a.cargo_resumido AS cargo,
            COUNT(*) AS dias_asistidos
        FROM asistencia a
        WHERE DATE(a.fecha_base) BETWEEN :start AND :end
          AND a.entrada IS NOT NULL
          {extra_asist}
          {extra_cta_asist}
          AND (:cta = '' OR a.cuenta_area = :cta)
        GROUP BY a.rut_trabajador, a.nombre, a.apellido_paterno, a.apellido_materno,
                 a.nombre_recinto, a.cuenta_area, a.cargo_resumido
    ),
    inas AS (
        SELECT
            REPLACE(REPLACE(i.DNI,'.',''),'-','') AS dni_norm,
            at.cuenta_area AS cuenta_area,
            COUNT(*) AS dias_inasistentes
        FROM inasistencias i
        JOIN asignacion_turnos at
          ON i.uid_inasistencia = at.uid_rut_dia_obra
        WHERE DATE(i.fecha_inasistencia) BETWEEN :start AND :end
          {extra_inas}
          {extra_cta_inas}
          AND (:cta = '' OR at.cuenta_area = :cta)
        GROUP BY dni_norm, at.cuenta_area
    )
    SELECT
        a.rut,
        a.nombre_completo AS nombre,
        a.recinto,
        a.cuenta_area,
        a.cargo,
        COALESCE(a.dias_asistidos,0)      AS dias_asistidos,
        COALESCE(i.dias_inasistentes,0)   AS dias_inasistentes,
        (COALESCE(a.dias_asistidos,0) + COALESCE(i.dias_inasistentes,0)) AS total_dias,
        ROUND(COALESCE(a.dias_asistidos,0) / NULLIF((COALESCE(a.dias_asistidos,0) + COALESCE(i.dias_inasistentes,0)),0) * 100, 1) AS pct_asistencia
    FROM asist a
    LEFT JOIN inas i
      ON REPLACE(REPLACE(a.rut,'.',''),'-','') = i.dni_norm
     AND a.cuenta_area = i.cuenta_area
    ORDER BY a.recinto, a.cuenta_area, a.nombre_completo
    """

    params = {"start": start, "end": end, "cta": cuenta_area,
              **p_asist, **p_inas, **p_cta_asist, **p_cta_inas}

    rows = db.session.execute(text(sql), params).mappings().all()
    df = pd.DataFrame([dict(r) for r in rows])

    buf = BytesIO()
    fname = f"nomina_{start}_a_{end}"

    try:
        import xlsxwriter; engine = "xlsxwriter"
    except Exception:
        try: import openpyxl; engine = "openpyxl"
        except Exception: engine = None

    if engine:
        with pd.ExcelWriter(buf, engine=engine) as writer:
            df.to_excel(writer, index=False, sheet_name="Nómina")
        buf.seek(0)
        return send_file(buf, as_attachment=True,
                         download_name=f"{fname}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    df.to_csv(buf, index=False, encoding="utf-8-sig"); buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{fname}.csv", mimetype="text/csv")



#================== DASHBOARD rotacion ===================
# ---- helpers fecha ----
def parse_ddmmyyyy(s: str):
    if not s:
        return None
    s = s.strip()
    for fmt in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None

def month_bounds(y: int, m: int):
    start = date(y, m, 1)
    end = date(y, m, monthrange(y, m)[1])
    return start, end

def headcount_on(when: date, cuenta: str | None = None, cargo: str | None = None,
                 allowed_ctas: set[str] | None = None) -> int:
    """
    Headcount (activos) en una fecha:
    FECHA_CTTO <= when y (FECHA_TERMINO is NULL o > when)
    Filtra por centro_costo_area (cuenta) y/o allowed_ctas si corresponde.
    """
    q = db.session.query(func.count(Desvinculacion.id)).filter(
        and_(
            Desvinculacion.FECHA_CTTO <= when,
            or_(Desvinculacion.FECHA_TERMINO.is_(None),
                Desvinculacion.FECHA_TERMINO > when),
        )
    )

    # Filtro por cuenta explícita
    if cuenta:
        q = q.filter(Desvinculacion.centro_costo_area == cuenta)
    # Si no hay cuenta explícita y el usuario tiene restricción de cuentas
    elif allowed_ctas is not None:
        if not allowed_ctas:
            return 0
        q = q.filter(Desvinculacion.centro_costo_area.in_(list(allowed_ctas)))

    if cargo:
        q = q.filter(Desvinculacion.CARGO == cargo)

    return q.scalar() or 0


@bp.route("/rotacion", methods=["GET"])
@login_required
def rotacion_filtros():
    # -- Rango por defecto: mes en curso --
    today = date.today()
    ini_def = date(today.year, today.month, 1)
    fin_def = (date(today.year + (today.month // 12), (today.month % 12) + 1, 1)
               - timedelta(days=1))

    ini = parse_ddmmyyyy(request.args.get("desde")) or ini_def
    fin = parse_ddmmyyyy(request.args.get("hasta")) or fin_def

    # ahora 'cuenta' = centro_costo_area
    cuenta = (request.args.get("cuenta") or "").strip() or None
    cargo  = (request.args.get("cargo") or "").strip() or None

    # --- PERMISOS: set de cuentas visibles para el usuario ---
    allowed_ctas = _allowed_cuentas_flat_for_current_user()  # None / set() / {'PGC','BRF',...}

    # Si el usuario seleccionó una cuenta explícita, validar que tenga acceso
    if cuenta and (allowed_ctas is not None) and (cuenta not in allowed_ctas):
        abort(403)

    # ---------- Opciones de SELECTS (cuentas y cargos) ----------
    # Cuentas visibles (desde centro_costo_area)
    cq = (db.session.query(Desvinculacion.centro_costo_area)
          .filter(Desvinculacion.centro_costo_area.isnot(None))
          .distinct()
          .order_by(Desvinculacion.centro_costo_area.asc()))

    if allowed_ctas is not None:
        if not allowed_ctas:
            cuentas_opts = []
        else:
            cq = cq.filter(Desvinculacion.centro_costo_area.in_(list(allowed_ctas)))
            cuentas_opts = [r[0] for r in cq.all()]
    else:
        cuentas_opts = [r[0] for r in cq.all()]

    # Cargos
    gq = (db.session.query(Desvinculacion.CARGO)
          .filter(Desvinculacion.CARGO.isnot(None))
          .distinct()
          .order_by(Desvinculacion.CARGO.asc()))
    cargos_opts = [r[0] for r in gq.all()]

    # ---------- Desvinculaciones del período (d) ----------
    dq = db.session.query(func.count(Desvinculacion.id)).filter(
        Desvinculacion.FECHA_TERMINO >= ini,
        Desvinculacion.FECHA_TERMINO <= fin,
    )

    if cuenta:
        dq = dq.filter(Desvinculacion.centro_costo_area == cuenta)

    # Aplicar permisos solo si hay restricción (allowed_ctas != None)
    if allowed_ctas is not None:
        if not allowed_ctas:
            d = 0
        else:
            # Si no hay cuenta seleccionada, filtrar por el set permitido
            if not cuenta:
                dq = dq.filter(Desvinculacion.centro_costo_area.in_(list(allowed_ctas)))
            d = dq.scalar() or 0
    else:
        d = dq.scalar() or 0

    # ---------- Headcount y rotación ----------
    Ai = headcount_on(ini, cuenta, cargo, allowed_ctas=allowed_ctas)
    Af = headcount_on(fin, cuenta, cargo, allowed_ctas=allowed_ctas)
    avg_dot = (Ai + Af) / 2 if (Ai is not None and Af is not None) else None
    rotacion = round((d / avg_dot) * 100, 2) if avg_dot and avg_dot > 0 else None

    # ---------- Series por mes ----------
    labels_m, activos_m, desv_m, rot_m = [], [], [], []
    cur = date(ini.year, ini.month, 1)
    end_month = date(fin.year, fin.month, 1)

    while cur <= end_month:
        m_ini, m_fin = month_bounds(cur.year, cur.month)
        if m_fin < ini:
            cur = (m_fin + timedelta(days=1)).replace(day=1)
            continue
        if m_ini > fin:
            break

        m_from, m_to = max(m_ini, ini), min(m_fin, fin)

        md_q = db.session.query(func.count(Desvinculacion.id)).filter(
            Desvinculacion.FECHA_TERMINO >= m_from,
            Desvinculacion.FECHA_TERMINO <= m_to,
        )

        if cuenta:
            md_q = md_q.filter(Desvinculacion.centro_costo_area == cuenta)

        if allowed_ctas is not None:
            if not allowed_ctas:
                md = 0
            else:
                if not cuenta:
                    md_q = md_q.filter(Desvinculacion.centro_costo_area.in_(list(allowed_ctas)))
                md = md_q.scalar() or 0
        else:
            md = md_q.scalar() or 0

        mAi = headcount_on(m_from, cuenta, cargo, allowed_ctas=allowed_ctas)
        mAf = headcount_on(m_to,   cuenta, cargo, allowed_ctas=allowed_ctas)
        mavg = (mAi + mAf) / 2 if (mAi + mAf) > 0 else 0
        mrot = (md / mavg) * 100 if mavg > 0 else 0

        labels_m.append(m_from.strftime("%b %Y"))
        activos_m.append(round(mavg, 0))
        desv_m.append(md)
        rot_m.append(round(mrot, 2))

        cur = date(cur.year + (cur.month // 12), (cur.month % 12) + 1, 1)

    # ---------- Top por cargo ----------
    base_cargo = db.session.query(
        Desvinculacion.CARGO,
        func.count(Desvinculacion.id).label("d")
    ).filter(
        Desvinculacion.FECHA_TERMINO >= ini,
        Desvinculacion.FECHA_TERMINO <= fin,
    )
    if cuenta:
        base_cargo = base_cargo.filter(Desvinculacion.centro_costo_area == cuenta)
    elif allowed_ctas is not None:
        if not allowed_ctas:
            base_cargo = base_cargo.filter(False)
        else:
            base_cargo = base_cargo.filter(Desvinculacion.centro_costo_area.in_(list(allowed_ctas)))

    base_cargo = (base_cargo
                  .group_by(Desvinculacion.CARGO)
                  .order_by(func.count(Desvinculacion.id).desc())
                  .limit(10).all())

    labels_cargo, desv_cargo, rot_cargo = [], [], []
    for c_name, d_cnt in base_cargo:
        cAi = headcount_on(ini, cuenta, c_name, allowed_ctas=allowed_ctas)
        cAf = headcount_on(fin, cuenta, c_name, allowed_ctas=allowed_ctas)
        cavg = (cAi + cAf) / 2 if (cAi + cAf) > 0 else 0
        crot = (d_cnt / cavg) * 100 if cavg > 0 else 0
        labels_cargo.append(c_name or "—")
        desv_cargo.append(int(d_cnt))
        rot_cargo.append(round(crot, 2))

    # ---------- Top por cuenta / área ----------
    base_cuenta = db.session.query(
        Desvinculacion.centro_costo_area,
        func.count(Desvinculacion.id).label("d")
    ).filter(
        Desvinculacion.FECHA_TERMINO >= ini,
        Desvinculacion.FECHA_TERMINO <= fin,
    )
    if cargo:
        base_cuenta = base_cuenta.filter(Desvinculacion.CARGO == cargo)
    if cuenta:
        base_cuenta = base_cuenta.filter(Desvinculacion.centro_costo_area == cuenta)
    elif allowed_ctas is not None:
        if not allowed_ctas:
            base_cuenta = base_cuenta.filter(False)
        else:
            base_cuenta = base_cuenta.filter(Desvinculacion.centro_costo_area.in_(list(allowed_ctas)))

    base_cuenta = (base_cuenta
                   .group_by(Desvinculacion.centro_costo_area)
                   .order_by(func.count(Desvinculacion.id).desc())
                   .limit(10).all())

    labels_cuenta, desv_cuenta, rot_cuenta = [], [], []
    for u_name, d_cnt in base_cuenta:
        uAi = headcount_on(ini, u_name, cargo=None, allowed_ctas=allowed_ctas)
        uAf = headcount_on(fin, u_name, cargo=None, allowed_ctas=allowed_ctas)
        uavg = (uAi + uAf) / 2 if (uAi + uAf) > 0 else 0
        urot = (d_cnt / uavg) * 100 if uavg > 0 else 0
        labels_cuenta.append(u_name or "—")
        desv_cuenta.append(int(d_cnt))
        rot_cuenta.append(round(urot, 2))

    return render_template(
        "dashboard/rotacion_filtros.html",
        ini=ini, fin=fin, cuenta=cuenta, cargo=cargo,
        cuentas_opts=cuentas_opts, cargos_opts=cargos_opts,
        d=d, Ai=Ai, Af=Af, avg_dot=avg_dot, rotacion=rotacion,
        labels_m=labels_m, activos_m=activos_m, desv_m=desv_m, rot_m=rot_m,
        labels_cargo=labels_cargo, desv_cargo=desv_cargo, rot_cargo=rot_cargo,
        labels_cuenta=labels_cuenta, desv_cuenta=desv_cuenta, rot_cuenta=rot_cuenta,
    )